from __future__ import annotations

import sqlite3
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from plistlib import InvalidFileException
from typing import Any

from openpyxl import load_workbook

from core.database import get_connection
from core.person_dictionary import GENDER_MAP, IDENTITY_MAP, PRODUCTION_GROUP_CLASSIFY_MAP, SPECIALIZE_CLASSIFY_MAP


class PersonService:
    """人员数据业务服务。"""

    # =========================
    # 基础工具
    # =========================

    @staticmethod
    def _now() -> str:
        """返回当前本地时间。"""
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    @staticmethod
    def _row_to_person(row: sqlite3.Row) -> dict[str, Any]:
        """
        SQLite snake_case
        转换为前端 camelCase。
        """

        return {
            "id": row["id"],
            "name": row["name"],
            "department": row["department"],
            "jobTitle": row["job_title"],
            "identity": row["identity"],
            "specializeClassify": row["specialize_classify"],
            "education": row["education"],
            "gender": row["gender"],
            "productionGroupClassify": row["production_group_classify"],
            "createTime": row["created_at"],
        }

    @staticmethod
    def _parse_required_int(value: Any, field_name: str) -> int:
        """
        将必填枚举字段转换为整数。

        支持：
        1
        "1"

        不允许：
        None
        ""
        True / False
        非数字字符串
        """

        if value is None or value == "":
            raise ValueError(f"{field_name}不能为空")

        # bool 是 int 的子类，需要单独排除
        if isinstance(value, bool):
            raise ValueError(f"{field_name}格式错误")

        try:
            result = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{field_name}格式错误") from exc

        return result

    @staticmethod
    def _parse_optional_int(value: Any, field_name: str) -> int | None:
        """
        解析可为空的整数枚举字段。
        """

        if value is None or value == "":
            return None

        if isinstance(value, bool):
            raise ValueError(f"{field_name}格式错误")

        try:
            return int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{field_name}格式错误") from exc

    @staticmethod
    def _escape_like(value: str) -> str:
        """
        转义 LIKE 中的特殊字符。

        防止用户搜索 % 或 _ 时被当成通配符。
        """

        return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")

    @staticmethod
    def _excel_text(value: Any) -> str:
        """Excel 单元格统一转文本。"""

        if value is None:
            return ""

        return str(value).strip()

    @staticmethod
    def _parse_excel_option(
        value: Any, field_name: str, mapping: dict[str, int], *, required: bool = True
    ) -> int | None:
        text = PersonService._excel_text(value)

        if not text:
            if required:
                raise ValueError(f"{field_name}不能为空")

            return None

        result = mapping.get(text)

        if result is None:
            raise ValueError(f"{field_name}“{text}”不存在")

        return result

    @staticmethod
    def _is_empty_excel_row(row: tuple[Any, ...]) -> bool:
        """
        判断 Excel 数据行是否为空。

        None、空字符串、纯空格都视为空。
        """

        return not any(value is not None and str(value).strip() for value in row)

    def _prepare_import_rows(self, worksheet):
        """
        初始化人员导入 Excel。

        模板约定：
        - 第1行：标题/说明
        - 第2行：字段表头
        - 第3行开始：人员数据

        返回：
        - rows：从第3行开始的数据迭代器
        - column_index：字段名 -> Excel列索引
        """

        rows = worksheet.iter_rows(values_only=True)

        # 第1行：标题/说明
        next(rows, None)

        # 第2行：字段表头
        header_row = next(rows, None)

        if not header_row:
            raise ValueError("Excel 缺少字段表头")

        headers = [self._excel_text(value) for value in header_row]

        required_headers = ["姓名11", "部门", "职名", "身份", "专业分类", "文化（报表用）", "性别", "生产组分类（1.普速铁路房建设备巡检维修人员；2.高速铁路房建设备巡检维修人员；3.行车公寓人员）"]

        missing_headers = [header for header in required_headers if header not in headers]

        if missing_headers:
            raise ValueError("Excel 缺少字段：" + "、".join(missing_headers))

        column_index = {header: index for index, header in enumerate(headers) if header}

        return (rows, column_index)

    @contextmanager
    def _open_import_workbook(self, file_path: str):
        """
        打开人员导入 Excel。

        统一负责：
        - 文件是否存在
        - 文件格式校验
        - openpyxl 打开文件
        - workbook.close()
        """

        path = Path(file_path)

        if not path.exists():
            raise ValueError("Excel 文件不存在")

        if not path.is_file():
            raise ValueError("Excel 文件路径无效")

        if path.suffix.lower() != ".xlsx":
            raise ValueError("仅支持 .xlsx 文件")

        try:
            workbook = load_workbook(filename=path, read_only=True, data_only=True)

        except InvalidFileException as exc:
            raise ValueError("Excel 文件格式错误") from exc

        except Exception as exc:
            raise ValueError("Excel 文件读取失败") from exc

        try:
            yield (path, workbook)

        finally:
            workbook.close()

    # =========================
    # 人员数据校验
    # =========================

    def _validate_person(self, data: dict[str, Any]) -> dict[str, Any]:
        """
        校验新增/编辑人员数据。

        注意：
        前端表单校验负责用户体验；
        这里的后端校验负责数据完整性。
        """

        name = str(data.get("name") or "").strip()

        department = str(data.get("department") or "").strip()

        job_title = str(data.get("jobTitle") or "").strip()

        education_text = str(data.get("education") or "").strip()

        education = education_text if education_text else None

        if not name:
            raise ValueError("姓名不能为空")

        if len(name) > 50:
            raise ValueError("姓名不能超过50个字符")

        if not department:
            raise ValueError("部门不能为空")

        if not job_title:
            raise ValueError("职名不能为空")

        identity = self._parse_required_int(data.get("identity"), "身份类型")

        specialize_classify = self._parse_required_int(data.get("specializeClassify"), "专业分类")

        gender = self._parse_required_int(data.get("gender"), "性别")

        production_group_classify = self._parse_optional_int(data.get("productionGroupClassify"), "生产组类别")

        return {
            "name": name,
            "department": department,
            "job_title": job_title,
            "identity": identity,
            "specialize_classify": specialize_classify,
            "education": education,
            "gender": gender,
            "production_group_classify": production_group_classify,
        }

    # =========================
    # 分页查询
    # =========================

    def list(self, params: dict[str, Any] | None = None) -> dict[str, Any]:
        """
        查询人员列表。

        支持参数：

        {
            "keyword": "",
            "department": "",
            "identity": 1,
            "specializeClassify": 1,
            "education": "",
            "gender": 1,
            "productionGroupClassify": 1,
            "page": 1,
            "pageSize": 20
        }
        """

        params = params or {}

        keyword = str(params.get("keyword") or "").strip()

        department = str(params.get("department") or "").strip()

        education = str(params.get("education") or "").strip()

        # =========================
        # 枚举筛选参数
        # =========================

        identity = self._parse_optional_int(params.get("identity"), "身份类型")

        specialize_classify = self._parse_optional_int(params.get("specializeClassify"), "专业分类")

        gender = self._parse_optional_int(params.get("gender"), "性别")

        production_group_classify = self._parse_optional_int(params.get("productionGroupClassify"), "生产组类别")

        # =========================
        # 分页
        # =========================

        try:
            page = int(params.get("page", 1))
        except (TypeError, ValueError):
            page = 1

        try:
            page_size = int(params.get("pageSize", 20))
        except (TypeError, ValueError):
            page_size = 20

        page = max(page, 1)

        # 最多一次查询100条
        page_size = max(1, min(page_size, 100))

        offset = (page - 1) * page_size

        # =========================
        # 构造 WHERE
        # =========================

        conditions: list[str] = []
        args: list[Any] = []

        if keyword:
            escaped = self._escape_like(keyword)

            like_keyword = f"%{escaped}%"

            conditions.append(
                """
                (
                    name LIKE ? ESCAPE '\\'
                    OR job_title LIKE ? ESCAPE '\\'
                )
                """
            )

            args.extend([like_keyword, like_keyword])

        if department:
            conditions.append("department = ?")

            args.append(department)

        if identity is not None:
            conditions.append("identity = ?")

            args.append(identity)

        if specialize_classify is not None:
            conditions.append("specialize_classify = ?")

            args.append(specialize_classify)

        if education:
            conditions.append("education = ?")

            args.append(education)

        if gender is not None:
            conditions.append("gender = ?")

            args.append(gender)

        if production_group_classify is not None:
            conditions.append("production_group_classify = ?")

            args.append(production_group_classify)

        where_sql = ""

        if conditions:
            where_sql = "WHERE " + " AND ".join(conditions)

        # =========================
        # 查询 SQLite
        # =========================

        with get_connection() as conn:
            # 总数
            count_sql = f"""
                SELECT
                    COUNT(*) AS total
                FROM person
                {where_sql}
            """

            total_row = conn.execute(count_sql, args).fetchone()

            total = total_row["total"] if total_row else 0

            # 当前页
            query_sql = f"""
                SELECT
                    id,
                    name,
                    department,
                    job_title,
                    identity,
                    specialize_classify,
                    education,
                    gender,
                    production_group_classify,
                    created_at
                FROM person
                {where_sql}
                ORDER BY id DESC
                LIMIT ?
                OFFSET ?
            """

            rows = conn.execute(query_sql, [*args, page_size, offset]).fetchall()

        return {
            "items": [self._row_to_person(row) for row in rows],
            "total": total,
            "page": page,
            "pageSize": page_size,
        }

    # =========================
    # 新增人员
    # =========================

    def create(self, data: dict[str, Any]) -> dict[str, Any]:
        """
        新增人员。
        """

        person = self._validate_person(data)

        now = self._now()

        with get_connection() as conn:
            cursor = conn.execute(
                """
                INSERT INTO person (
                    name,
                    department,
                    job_title,
                    identity,
                    specialize_classify,
                    education,
                    gender,
                    production_group_classify,
                    created_at,
                    updated_at
                )
                VALUES (
                    ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?
                )
                """,
                (
                    person["name"],
                    person["department"],
                    person["job_title"],
                    person["identity"],
                    person["specialize_classify"],
                    person["education"],
                    person["gender"],
                    person["production_group_classify"],
                    now,
                    now,
                ),
            )

            person_id = cursor.lastrowid

        return {"id": person_id}

    # =========================
    # 修改人员
    # =========================

    def update(self, person_id: int, data: dict[str, Any]) -> None:
        """
        修改人员。
        """

        person = self._validate_person(data)

        now = self._now()

        with get_connection() as conn:
            cursor = conn.execute(
                """
                UPDATE person
                SET
                    name = ?,
                    department = ?,
                    job_title = ?,
                    identity = ?,
                    specialize_classify = ?,
                    education = ?,
                    gender = ?,
                    production_group_classify = ?,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    person["name"],
                    person["department"],
                    person["job_title"],
                    person["identity"],
                    person["specialize_classify"],
                    person["education"],
                    person["gender"],
                    person["production_group_classify"],
                    now,
                    person_id,
                ),
            )

            if cursor.rowcount == 0:
                raise ValueError("人员不存在")

    # =========================
    # 删除人员
    # =========================

    def delete(self, person_id: int) -> None:
        """
        删除人员。
        """

        with get_connection() as conn:
            cursor = conn.execute(
                """
                DELETE FROM person
                WHERE id = ?
                """,
                (person_id,),
            )

            if cursor.rowcount == 0:
                raise ValueError("人员不存在")

    def preview_import_file(self, file_path: str) -> dict[str, Any]:
        """
        解析人员 Excel，
        只生成预览，不写数据库。
        """

        with self._open_import_workbook(file_path) as (path, workbook):
            worksheet = workbook.active

            rows, column_index = self._prepare_import_rows(worksheet)

            preview: list[dict[str, Any]] = []

            errors: list[dict[str, Any]] = []

            total = 0
            success_count = 0

            for excel_row_number, row in enumerate(rows, start=3):
                # 跳过完全空白行
                if self._is_empty_excel_row(row):
                    continue

                total += 1

                try:
                    validated = self._parse_excel_person_row(row, column_index)

                    success_count += 1

                    # 最多返回100条给前端预览
                    if len(preview) < 100:
                        preview.append(
                            {
                                "row": excel_row_number,
                                "name": validated["name"],
                                "department": validated["department"],
                                "jobTitle": validated["job_title"],
                                "identity": validated["identity"],
                                "specializeClassify": validated["specialize_classify"],
                                "education": validated["education"],
                                "gender": validated["gender"],
                                "productionGroupClassify": validated["production_group_classify"],
                            }
                        )

                except ValueError as exc:
                    errors.append({"row": excel_row_number, "message": str(exc)})

            return {
                "fileName": path.name,
                "total": total,
                "successCount": success_count,
                "failureCount": len(errors),
                "preview": preview,
                "previewLimit": 100,
                "errors": errors,
            }

    def _parse_excel_person_row(self, row: tuple[Any, ...], column_index: dict[str, int]) -> dict[str, Any]:
        """
        将 Excel 一行转换并校验为人员数据。

        返回值使用数据库字段 snake_case：
        {
            "name": ...,
            "department": ...,
            "job_title": ...,
            ...
        }
        """

        def cell(name: str) -> Any:
            index = column_index[name]

            if index >= len(row):
                return None

            return row[index]

        person = {
            "name": self._excel_text(cell("姓名11")),
            "department": self._excel_text(cell("部门")),
            "jobTitle": self._excel_text(cell("职名")),
            "identity": self._parse_excel_option(cell("身份"), "身份", IDENTITY_MAP),
            "specializeClassify": self._parse_excel_option(cell("专业分类"), "专业分类", SPECIALIZE_CLASSIFY_MAP),
            "education": self._excel_text(cell("文化（报表用）")),
            "gender": self._parse_excel_option(cell("性别"), "性别", GENDER_MAP),
            "productionGroupClassify": self._parse_excel_option(
                cell("生产组分类（1.普速铁路房建设备巡检维修人员；2.高速铁路房建设备巡检维修人员；3.行车公寓人员）"), "生产组分类", PRODUCTION_GROUP_CLASSIFY_MAP, required=False
            ),
        }

        # 继续使用 CRUD 同一套业务校验
        return self._validate_person(person)

    def import_file(self, file_path: str) -> dict[str, Any]:
        """
        正式导入人员 Excel。
        """

        with self._open_import_workbook(file_path) as (_path, workbook):
            worksheet = workbook.active

            rows, column_index = self._prepare_import_rows(worksheet)

            total = 0
            success_count = 0

            errors: list[dict[str, Any]] = []

            batch_size = 500

            batch: list[tuple[Any, ...]] = []

            now = self._now()

            insert_sql = """
                INSERT INTO person (
                    name,
                    department,
                    job_title,
                    identity,
                    specialize_classify,
                    education,
                    gender,
                    production_group_classify,
                    created_at,
                    updated_at
                )
                VALUES (
                    ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?
                )
            """

            with get_connection() as conn:
                for excel_row_number, row in enumerate(rows, start=3):
                    # 跳过完全空白行
                    if self._is_empty_excel_row(row):
                        continue

                    total += 1

                    try:
                        person = self._parse_excel_person_row(row, column_index)

                        batch.append(
                            (
                                person["name"],
                                person["department"],
                                person["job_title"],
                                person["identity"],
                                person["specialize_classify"],
                                person["education"],
                                person["gender"],
                                person["production_group_classify"],
                                now,
                                now,
                            )
                        )

                        success_count += 1

                        if len(batch) >= batch_size:
                            conn.executemany(insert_sql, batch)

                            batch.clear()

                    except ValueError as exc:
                        errors.append({"row": excel_row_number, "message": str(exc)})

                if batch:
                    conn.executemany(insert_sql, batch)

            return {"total": total, "successCount": success_count, "failureCount": len(errors), "errors": errors}
