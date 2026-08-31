from __future__ import annotations

from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path
from typing import Any

import xlrd
from openpyxl import load_workbook

from core.database import get_connection
from core.problem_dictionary import (
    CHECK_TYPE_MAP,
    CHECK_WAY_MAP,
    PROBLEM_CLASSIFY_MAP,
    PROBLEM_LABEL_MAP,
    RISK_LEVEL_MAP,
    STATUS_MAP,
    YES_NO_MAP,
)

PROBLEM_REQUIRED_HEADERS = [
    "状态",
    "是否分解",
    "是否修订",
    "是否考核",
    "检查部门",
    "检查人员",
    "检查类型",
    "检查方式",
    "检查开始时间",
    "检查结束时间",
    "提交时间",
    "问题项点",
    "后果明示",
    "风险类型",
    "风险等级",
    "问题归类",
    "问题分项",
    "问题加分",
    "问题类型",
    "是否红线",
    "发现地点",
    "落实部门",
    "责任部门",
    "其他归类",
    "跨单位",
    "业务指导",
    "路外问题",
    "问题描述",
    "整改时限",
]

PROBLEM_OPTIONAL_HEADERS = [
    "整改要求",
    "整改人",
    "整改时间",
    "整改描述",
    "问题责任人",
    "问题原因",
    "销号人",
    "销号评价",
    "销号时间",
]

PROBLEM_INSERT_SQL = """
    INSERT INTO check_problem (
        status,
        decomposed,
        revised,
        assessed,
        check_department,
        check_person,
        check_type,
        check_way,
        check_start_time,
        check_end_time,
        submit_time,
        problem_item,
        consequences,
        risk_type,
        risk_level,
        problem_classify,
        problem_label,
        problem_extra_points,
        problem_type,
        red_line,
        problem_position,
        implement_department,
        duty_department,
        other_classify,
        cross_unit,
        business_guidance,
        outside,
        problem_description,
        deadline,
        rectification_requirements,
        rectification_person,
        rectification_time,
        rectification_description,
        responsible_person,
        reason,
        close_issue_person,
        close_issue_evaluate,
        close_issue_time
    )
    VALUES (
        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
        ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
        ?, ?, ?, ?, ?, ?, ?, ?
    )
"""

class ProblemService:
    def list(self, params: dict[str, Any]) -> dict[str, Any]:
        page = max(int(params.get("page") or 1), 1)
        page_size = min(max(int(params.get("pageSize") or 20), 1), 100)
        offset = (page - 1) * page_size

        conditions: list[str] = []
        values: list[Any] = []

        keyword = str(params.get("keyword") or "").strip()

        if keyword:
            keyword_value = f"%{keyword}%"

            conditions.append(
                "("
                "problem_description LIKE ? OR "
                "problem_item LIKE ? OR "
                "check_person LIKE ? OR "
                "check_department LIKE ? OR "
                "implement_department LIKE ? OR "
                "duty_department LIKE ? OR "
                "problem_position LIKE ?"
                ")"
            )

            values.extend([keyword_value] * 7)

        self._add_equal_condition(conditions, values, "status", params.get("status"))
        self._add_equal_condition(conditions, values, "check_type", params.get("checkType"))
        self._add_equal_condition(conditions, values, "check_way", params.get("checkWay"))
        self._add_equal_condition(conditions, values, "risk_level", params.get("riskLevel"))
        self._add_equal_condition(conditions, values, "problem_classify", params.get("problemClassify"))
        self._add_equal_condition(conditions, values, "problem_label", params.get("problemLabel"))
        self._add_equal_condition(conditions, values, "red_line", params.get("redLine"))

        check_department = str(params.get("checkDepartment") or "").strip()

        if check_department:
            conditions.append("check_department LIKE ?")
            values.append(f"%{check_department}%")

        check_person = str(params.get("checkPerson") or "").strip()

        if check_person:
            conditions.append("check_person LIKE ?")
            values.append(f"%{check_person}%")

        where_sql = ""

        if conditions:
            where_sql = " WHERE " + " AND ".join(conditions)

        count_sql = f"SELECT COUNT(*) FROM check_problem{where_sql}"

        list_sql = f"""
            SELECT *
            FROM check_problem
            {where_sql}
            ORDER BY submit_time DESC, id DESC
            LIMIT ? OFFSET ?
        """

        with get_connection() as conn:
            total = conn.execute(count_sql, values).fetchone()[0]

            rows = conn.execute(list_sql, [*values, page_size, offset]).fetchall()

        return {
            "items": [self._row_to_problem(row) for row in rows],
            "total": total,
            "page": page,
            "pageSize": page_size,
        }

    @staticmethod
    def _add_equal_condition(conditions: list[str], values: list[Any], column: str, value: Any) -> None:
        if value is None:
            return

        conditions.append(f"{column} = ?")
        values.append(value)

    @staticmethod
    def _row_to_problem(row) -> dict[str, Any]:
        return {
            "id": row["id"],
            "status": row["status"],
            "decomposed": row["decomposed"],
            "revised": row["revised"],
            "assessed": row["assessed"],
            "checkDepartment": row["check_department"],
            "checkPerson": row["check_person"],
            "checkType": row["check_type"],
            "checkWay": row["check_way"],
            "checkStartTime": row["check_start_time"],
            "checkEndTime": row["check_end_time"],
            "submitTime": row["submit_time"],
            "problemItem": row["problem_item"],
            "consequences": row["consequences"],
            "riskType": row["risk_type"],
            "riskLevel": row["risk_level"],
            "problemClassify": row["problem_classify"],
            "problemLabel": row["problem_label"],
            "problemExtraPoints": row["problem_extra_points"],
            "problemType": row["problem_type"],
            "redLine": row["red_line"],
            "problemPosition": row["problem_position"],
            "implementDepartment": row["implement_department"],
            "dutyDepartment": row["duty_department"],
            "otherClassify": row["other_classify"],
            "crossUnit": row["cross_unit"],
            "businessGuidance": row["business_guidance"],
            "outside": row["outside"],
            "problemDescription": row["problem_description"],
            "deadline": row["deadline"],
            "rectificationRequirements": row["rectification_requirements"],
            "rectificationPerson": row["rectification_person"],
            "rectificationTime": row["rectification_time"],
            "rectificationDescription": row["rectification_description"],
            "responsiblePerson": row["responsible_person"],
            "reason": row["reason"],
            "closeIssuePerson": row["close_issue_person"],
            "closeIssueEvaluate": row["close_issue_evaluate"],
            "closeIssueTime": row["close_issue_time"],
        }

    @contextmanager
    def _open_import_rows(self, file_path: str):
        path = Path(file_path)

        if not path.exists() or not path.is_file():
            raise ValueError("Excel 文件不存在")

        suffix = path.suffix.lower()

        if suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)

            try:
                worksheet = workbook.active
                yield iter(worksheet.iter_rows(values_only=True))
            finally:
                workbook.close()

            return

        if suffix == ".xls":
            workbook = xlrd.open_workbook(path)
            worksheet = workbook.sheet_by_index(0)

            yield self._iter_xls_rows(worksheet, workbook.datemode)
            return

        raise ValueError("仅支持 .xls 和 .xlsx 文件")

    @staticmethod
    def _iter_xls_rows(worksheet, datemode: int):
        for row_index in range(worksheet.nrows):
            row = []

            for column_index in range(worksheet.ncols):
                cell = worksheet.cell(row_index, column_index)
                value = cell.value

                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, datemode)

                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    value = None

                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(value).is_integer():
                    value = int(value)

                row.append(value)

            yield tuple(row)

    def _prepare_import_rows(self, rows):
        rows = iter(rows)

        # 第1行：标题
        next(rows, None)

        # 第2行：字段名
        header_row = next(rows, None)

        if header_row is None:
            raise ValueError("Excel 第2行不存在字段名")

        column_index: dict[str, int] = {}

        for index, value in enumerate(header_row):
            if value is None:
                continue

            header = str(value).strip()

            if header:
                column_index[header] = index

        missing_headers = [header for header in PROBLEM_REQUIRED_HEADERS if header not in column_index]

        if missing_headers:
            raise ValueError(f"缺少必填列：{', '.join(missing_headers)}")

        return rows, column_index

    @staticmethod
    def _is_empty_excel_row(row) -> bool:
        return not any(value is not None and str(value).strip() for value in row)

    @staticmethod
    def _get_excel_value(row, column_index: dict[str, int], header: str):
        index = column_index.get(header)

        if index is None or index >= len(row):
            return None

        return row[index]

    @staticmethod
    def _parse_option(value: Any, option_map: dict[str, int], field_name: str) -> int:
        if value is None or str(value).strip() == "":
            raise ValueError(f"{field_name}不能为空")

        # 兼容 Excel 直接填写 0 / 1 / 2
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            if float(value).is_integer():
                option_value = int(value)

                if option_value in option_map.values():
                    return option_value

        text = str(value).strip()

        if text in option_map:
            return option_map[text]

        if text.isdigit():
            option_value = int(text)

            if option_value in option_map.values():
                return option_value

        allowed = "、".join(option_map.keys())

        raise ValueError(f"{field_name}值无效：{text}，可选值：{allowed}")

    @staticmethod
    def _parse_required_text(value: Any, field_name: str) -> str:
        if value is None:
            raise ValueError(f"{field_name}不能为空")

        text = str(value).strip()

        if not text:
            raise ValueError(f"{field_name}不能为空")

        return text

    @staticmethod
    def _parse_optional_text(value: Any) -> str | None:
        if value is None:
            return None

        text = str(value).strip()

        return text if text else None

    @staticmethod
    def _parse_datetime(value: Any, field_name: str, required: bool = True) -> str | None:
        if value is None or str(value).strip() == "":
            if required:
                raise ValueError(f"{field_name}不能为空")

            return None

        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")

        if isinstance(value, date):
            return f"{value:%Y-%m-%d} 00:00:00"

        text = str(value).strip()

        formats = [
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y/%m/%d %H:%M:%S",
            "%Y/%m/%d %H:%M",
            "%Y-%m-%d",
            "%Y/%m/%d",
        ]

        for date_format in formats:
            try:
                parsed = datetime.strptime(text, date_format)
                return parsed.strftime("%Y-%m-%d %H:%M:%S")
            except ValueError:
                continue

        raise ValueError(f"{field_name}时间格式无效：{text}")

    @staticmethod
    def _parse_date(value: Any, field_name: str) -> str:
        if value is None or str(value).strip() == "":
            raise ValueError(f"{field_name}不能为空")

        if isinstance(value, datetime):
            return value.date().isoformat()

        if isinstance(value, date):
            return value.isoformat()

        text = str(value).strip()

        formats = [
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%Y-%m-%d %H:%M:%S",
            "%Y-%m-%d %H:%M",
            "%Y/%m/%d %H:%M:%S",
            "%Y/%m/%d %H:%M",
        ]

        for date_format in formats:
            try:
                return datetime.strptime(text, date_format).date().isoformat()
            except ValueError:
                continue

        raise ValueError(f"{field_name}日期格式无效：{text}")

    def _parse_excel_problem_row(self, row, column_index: dict[str, int]) -> dict[str, Any]:
        get = lambda header: self._get_excel_value(row, column_index, header)

        return {
            "status": self._parse_option(get("状态"), STATUS_MAP, "状态"),
            "decomposed": self._parse_option(get("是否分解"), YES_NO_MAP, "是否分解"),
            "revised": self._parse_option(get("是否修订"), YES_NO_MAP, "是否修订"),
            "assessed": self._parse_option(get("是否考核"), YES_NO_MAP, "是否考核"),
            "checkDepartment": self._parse_required_text(get("检查部门"), "检查部门"),
            "checkPerson": self._parse_required_text(get("检查人员"), "检查人员"),
            "checkType": self._parse_option(get("检查类型"), CHECK_TYPE_MAP, "检查类型"),
            "checkWay": self._parse_option(get("检查方式"), CHECK_WAY_MAP, "检查方式"),
            "checkStartTime": self._parse_datetime(get("检查开始时间"), "检查开始时间"),
            "checkEndTime": self._parse_datetime(get("检查结束时间"), "检查结束时间"),
            "submitTime": self._parse_datetime(get("提交时间"), "提交时间"),
            "problemItem": self._parse_required_text(get("问题项点"), "问题项点"),
            "consequences": self._parse_required_text(get("后果明示"), "后果明示"),
            "riskType": self._parse_required_text(get("风险类型"), "风险类型"),
            "riskLevel": self._parse_option(get("风险等级"), RISK_LEVEL_MAP, "风险等级"),
            "problemClassify": self._parse_option(get("问题归类"), PROBLEM_CLASSIFY_MAP, "问题归类"),
            "problemLabel": self._parse_option(get("问题分项"), PROBLEM_LABEL_MAP, "问题分项"),
            "problemExtraPoints": self._parse_required_text(get("问题加分"), "问题加分"),
            "problemType": self._parse_required_text(get("问题类型"), "问题类型"),
            "redLine": self._parse_option(get("是否红线"), YES_NO_MAP, "是否红线"),
            "problemPosition": self._parse_required_text(get("发现地点"), "发现地点"),
            "implementDepartment": self._parse_required_text(get("落实部门"), "落实部门"),
            "dutyDepartment": self._parse_required_text(get("责任部门"), "责任部门"),
            "otherClassify": self._parse_required_text(get("其他归类"), "其他归类"),
            "crossUnit": self._parse_option(get("跨单位"), YES_NO_MAP, "跨单位"),
            "businessGuidance": self._parse_option(get("业务指导"), YES_NO_MAP, "业务指导"),
            "outside": self._parse_option(get("路外问题"), YES_NO_MAP, "路外问题"),
            "problemDescription": self._parse_required_text(get("问题描述"), "问题描述"),
            "deadline": self._parse_date(get("整改时限"), "整改时限"),
            "rectificationRequirements": self._parse_optional_text(get("整改要求")),
            "rectificationPerson": self._parse_optional_text(get("整改人")),
            "rectificationTime": self._parse_datetime(get("整改时间"), "整改时间", required=False),
            "rectificationDescription": self._parse_optional_text(get("整改描述")),
            "responsiblePerson": self._parse_optional_text(get("问题责任人")),
            "reason": self._parse_optional_text(get("问题原因")),
            "closeIssuePerson": self._parse_optional_text(get("销号人")),
            "closeIssueEvaluate": self._parse_optional_text(get("销号评价")),
            "closeIssueTime": self._parse_datetime(get("销号时间"), "销号时间", required=False),
        }

    @staticmethod
    def _problem_to_insert_values(problem: dict[str, Any]) -> tuple[Any, ...]:
        return (
            problem["status"],
            problem["decomposed"],
            problem["revised"],
            problem["assessed"],
            problem["checkDepartment"],
            problem["checkPerson"],
            problem["checkType"],
            problem["checkWay"],
            problem["checkStartTime"],
            problem["checkEndTime"],
            problem["submitTime"],
            problem["problemItem"],
            problem["consequences"],
            problem["riskType"],
            problem["riskLevel"],
            problem["problemClassify"],
            problem["problemLabel"],
            problem["problemExtraPoints"],
            problem["problemType"],
            problem["redLine"],
            problem["problemPosition"],
            problem["implementDepartment"],
            problem["dutyDepartment"],
            problem["otherClassify"],
            problem["crossUnit"],
            problem["businessGuidance"],
            problem["outside"],
            problem["problemDescription"],
            problem["deadline"],
            problem["rectificationRequirements"],
            problem["rectificationPerson"],
            problem["rectificationTime"],
            problem["rectificationDescription"],
            problem["responsiblePerson"],
            problem["reason"],
            problem["closeIssuePerson"],
            problem["closeIssueEvaluate"],
            problem["closeIssueTime"],
        )

    def preview_import_file(self, file_path: str) -> dict[str, Any]:
        preview: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []

        total = 0
        success_count = 0
        failure_count = 0

        with self._open_import_rows(file_path) as excel_rows:
            rows, column_index = self._prepare_import_rows(excel_rows)

            for excel_row, row in enumerate(rows, start=3):
                if self._is_empty_excel_row(row):
                    continue

                total += 1

                try:
                    problem = self._parse_excel_problem_row(row, column_index)

                    success_count += 1

                    if len(preview) < 100:
                        preview.append({"row": excel_row, **problem})

                except ValueError as exc:
                    failure_count += 1
                    errors.append({"row": excel_row, "message": str(exc)})

        return {
            "fileName": Path(file_path).name,
            "total": total,
            "successCount": success_count,
            "failureCount": failure_count,
            "preview": preview,
            "previewLimit": 100,
            "errors": errors,
        }

    def import_file(self, file_path: str) -> dict[str, Any]:
        errors: list[dict[str, Any]] = []

        total = 0
        success_count = 0
        failure_count = 0

        batch: list[tuple[Any, ...]] = []

        with self._open_import_rows(file_path) as excel_rows:
            rows, column_index = self._prepare_import_rows(excel_rows)

            with get_connection() as conn:
                for excel_row, row in enumerate(rows, start=3):
                    if self._is_empty_excel_row(row):
                        continue

                    total += 1

                    try:
                        problem = self._parse_excel_problem_row(row, column_index)
                        batch.append(self._problem_to_insert_values(problem))

                        success_count += 1

                        if len(batch) >= 500:
                            conn.executemany(PROBLEM_INSERT_SQL, batch)
                            batch.clear()

                    except ValueError as exc:
                        failure_count += 1
                        errors.append({"row": excel_row, "message": str(exc)})

                if batch:
                    conn.executemany(PROBLEM_INSERT_SQL, batch)

        return {
            "total": total,
            "successCount": success_count,
            "failureCount": failure_count,
            "errors": errors,
        }